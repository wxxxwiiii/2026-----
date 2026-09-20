"""智能备课台：首页、AI 备课中心、教材资源库。"""
import io
import json
import os
import re
import time
import uuid
import difflib

from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify, send_file
from werkzeug.utils import secure_filename

from config import UPLOAD_DIR, TMP_DIR
from extensions import db
from models import (Resource, KbDoc, KbChunk, LessonPlan, SchoolClass, Student, StudentLearning,
                    LayeredAssignment, ClassSportBinding, ClassPeriod, LearningSheet, SPORT_ITEMS,
                    COURSE_NAMES, SPORT_VENUES, PROJECT_ITEMS)
from security import login_required, current_user_id
from services.ai import provider_options

lesson_prep_bp = Blueprint('lesson_prep', __name__, url_prefix='/lesson-prep')

RESOURCE_CATEGORIES = ['教案模板', '教学进度表', '教学课件（PPT）', '教学视频', '课标与政策文件', '试题与测评量表', '图书 / 教学用书', '图片素材', '通用文档', '其他']

GRADES = ['大一', '大二', '大三', '大四', '研究生', '通用']


@lesson_prep_bp.route('/')
@login_required
def index():
    return render_template('lesson_prep/index.html', active='lesson_prep',
                           providers=provider_options())


@lesson_prep_bp.route('/ai-center')
@login_required
def ai_center():
    return render_template('lesson_prep/ai_center.html', active='ai_center',
                           providers=provider_options())


@lesson_prep_bp.route('/resources')
@login_required
def resources():
    cat = request.args.get('category', '')
    q = Resource.query
    if cat:
        q = q.filter_by(category=cat)
    items = q.order_by(Resource.id.desc()).all()
    classes = _all_classes()
    class_map = {c.id: c.name for c in classes}
    # 进度表资源 → 已同步班级（双向溯源）
    resource_classes = {}
    progress_ids = [r.id for r in items if r.category == '教学进度表']
    if progress_ids:
        for p in ClassPeriod.query.filter(ClassPeriod.source_file_id.in_(progress_ids)).all():
            if p.source_file_id and p.class_id:
                resource_classes.setdefault(p.source_file_id, set()).add((p.class_id, class_map.get(p.class_id, '')))
    resource_classes = {rid: sorted(lst) for rid, lst in resource_classes.items()}
    return render_template('lesson_prep/resources.html', active='resources',
                           items=items, categories=RESOURCE_CATEGORIES, current_cat=cat,
                           classes=classes, course_names=COURSE_NAMES,
                           project_items=PROJECT_ITEMS, sport_venues=SPORT_VENUES,
                           resource_classes=resource_classes)


@lesson_prep_bp.route('/resources/add', methods=['POST'])
@login_required
def resource_add():
    name = request.form.get('name', '').strip()
    category = request.form.get('category', '其他')
    link = request.form.get('link', '').strip()
    note = request.form.get('note', '').strip()
    if not name:
        flash('请填写资源名称', 'error')
        return redirect(url_for('lesson_prep.resources'))

    file_path = ''
    saved_full = ''
    f = request.files.get('file')
    if f and f.filename:
        ext = os.path.splitext(secure_filename(f.filename))[1].lower()
        if len(ext) > 12:
            ext = ''
        filename = 'resource_%d%s' % (int(time.time() * 1000), ext)
        saved_full = os.path.join(UPLOAD_DIR, filename)
        f.save(saved_full)
        file_path = '/static/uploads/' + filename

    db.session.add(Resource(name=name, category=category, file_path=file_path, link=link, note=note))

    # 同步写入本地知识库，供统一检索
    doc = KbDoc(owner_id=current_user_id(), title=name, doc_type=category, source='local', url=link)
    db.session.add(doc)
    db.session.flush()
    chunk_text = name + ' ' + note
    if saved_full and os.path.splitext(saved_full)[1].lower() in ('.txt', '.md'):
        try:
            with open(saved_full, 'r', encoding='utf-8', errors='ignore') as fp:
                chunk_text += '\n' + fp.read()[:2000]
        except Exception:
            pass
    db.session.add(KbChunk(doc_id=doc.id, text=chunk_text, source_name='本地上传'))

    db.session.commit()
    flash('资源已添加', 'success')
    return redirect(url_for('lesson_prep.resources'))


@lesson_prep_bp.route('/resources/<int:rid>/delete', methods=['POST'])
@login_required
def resource_delete(rid):
    r = db.get_or_404(Resource, rid)
    db.session.delete(r)
    db.session.commit()
    flash('资源已删除', 'success')
    return redirect(url_for('lesson_prep.resources'))


# ============ 教学进度表上传导入 ============

PROGRESS_HEADERS = ['课次', '授课内容', '星期', '节次', '本次学时', '作业', '备注']
ALLOWED_PROGRESS_EXT = ('.xlsx', '.doc', '.docx')


def _cell_str(v):
    """把单元格值转成去空白字符串。"""
    if v is None:
        return ''
    return str(v).strip()


def _parse_order(v):
    """解析课次为整数，非数字返回 None。"""
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    s = _cell_str(v)
    if not s:
        return None
    try:
        return int(s)
    except (ValueError, TypeError):
        return None


def _build_progress_template():
    """生成标准教学进度表 Excel 模板，返回 BytesIO。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = '教学进度表'
    for i, h in enumerate(PROGRESS_HEADERS, start=1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = Font(bold=True)
    for i, w in enumerate((8, 40, 10, 10, 10, 24, 24), start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws2 = wb.create_sheet('使用说明')
    for i, line in enumerate([
        '教学进度表导入说明',
        '1. 一行代表一个课次，不要把多个课次写在同一行。',
        '2. 最多支持 50 条课次数据。',
        '3. 必填列：课次（数字，不能重复）、授课内容。',
        '4. 选填列：星期、节次、本次学时、作业、备注，可留空。',
        '5. 课程名称、授课项目、授课地点不写在表格内，上传时在网页选择。',
    ], start=1):
        ws2.cell(row=i, column=1, value=line)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _parse_progress_xlsx(file_stream):
    """解析标准 Excel 模板，返回 (rows, error, warning)。"""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(file_stream, read_only=True, data_only=True)
    except Exception:
        return None, '文件解析失败，请下载标准模板整理后重新上传', None
    raw = list(wb.active.iter_rows(values_only=True))
    if not raw:
        return None, '文件为空，请下载标准模板整理后重新上传', None
    header = [_cell_str(c) for c in raw[0]]
    col = {}
    for i, h in enumerate(header):
        for key, label in (('order', '课次'), ('content', '授课内容'), ('week', '星期'),
                           ('section', '节次'), ('hours', '本次学时'), ('homework', '作业'), ('note', '备注')):
            if h == label:
                col[key] = i
    if 'order' not in col or 'content' not in col:
        return None, '文件格式不对，缺少必填列「课次/授课内容」，请使用标准模板', None
    rows = []
    seen = set()
    for idx, r in enumerate(raw[1:], start=1):
        if not r or all(_cell_str(c) == '' for c in r):
            continue
        order = _parse_order(r[col['order']]) if len(r) > col['order'] else None
        if order is None:
            return None, '课次必须填写数字（第 %d 行数据）' % (idx + 1), None
        if order in seen:
            return None, '课次 %d 重复，请检查表格数据' % order, None
        seen.add(order)
        content = _cell_str(r[col['content']]) if len(r) > col['content'] else ''
        if not content:
            return None, '第 %d 行授课内容不能为空' % (idx + 1), None
        rows.append({
            'order_no': order, 'content': content,
            'week': _cell_str(r[col['week']]) if 'week' in col and len(r) > col['week'] else '',
            'section': _cell_str(r[col['section']]) if 'section' in col and len(r) > col['section'] else '',
            'hours': _cell_str(r[col['hours']]) if 'hours' in col and len(r) > col['hours'] else '',
            'homework': _cell_str(r[col['homework']]) if 'homework' in col and len(r) > col['homework'] else '',
            'note': _cell_str(r[col['note']]) if 'note' in col and len(r) > col['note'] else '',
        })
    if not rows:
        return None, '未解析到课次数据，请下载标准模板整理后重新上传', None
    if len(rows) > 50:
        return None, '最多支持 50 条课次数据，当前 %d 条' % len(rows), None
    return rows, None, None


def _parse_progress_docx(file_stream):
    """解析旧版 Word 教学进度表：周次→课次、讲授章节及内容→授课内容。返回 (rows, error, warning)。"""
    from docx import Document
    try:
        doc = Document(file_stream)
    except Exception:
        return None, '文件解析失败，请下载标准模板整理后重新上传', None
    if not doc.tables:
        return None, '未识别到表格，请下载标准模板整理后重新上传', None
    rows = []
    seen = set()
    for table in doc.tables:
        data = [[_cell_str(c.text) for c in row.cells] for row in table.rows]
        if not data:
            continue
        order_col = content_col = week_col = section_col = hours_col = homework_col = note_col = None
        for i, h in enumerate(data[0]):
            if order_col is None and ('周次' in h or '课次' in h):
                order_col = i
            elif content_col is None and ('讲授' in h or '内容' in h):
                content_col = i
            elif week_col is None and '星期' in h:
                week_col = i
            elif section_col is None and '节次' in h:
                section_col = i
            elif hours_col is None and ('学时' in h or '课时' in h):
                hours_col = i
            elif homework_col is None and '作业' in h:
                homework_col = i
            elif note_col is None and '备注' in h:
                note_col = i
        if order_col is None or content_col is None:
            continue
        for r in data[1:]:
            order = _parse_order(r[order_col]) if order_col < len(r) else None
            content = r[content_col] if content_col < len(r) else ''
            if order is None or not content or order in seen:
                continue
            seen.add(order)
            rows.append({
                'order_no': order, 'content': content,
                'week': r[week_col] if week_col is not None and week_col < len(r) else '',
                'section': r[section_col] if section_col is not None and section_col < len(r) else '',
                'hours': r[hours_col] if hours_col is not None and hours_col < len(r) else '',
                'homework': r[homework_col] if homework_col is not None and homework_col < len(r) else '',
                'note': r[note_col] if note_col is not None and note_col < len(r) else '',
            })
    if not rows:
        return None, '未识别到「周次/讲授章节及内容」列，请下载标准模板整理后重新上传', None
    if len(rows) > 50:
        return None, '最多支持 50 条课次数据，当前 %d 条' % len(rows), None
    return rows, None, '检测为旧版Word进度表，已完成智能解析，请仔细核对预览数据，建议后续使用标准Excel模板'


def _period_label(order_no, content):
    """生成课时短标题（供分层下拉展示）。"""
    first = (content or '').strip().splitlines()[0].strip() if content else ''
    if len(first) > 24:
        first = first[:24] + '…'
    return ('第%d课 %s' % (order_no, first)) if first else ('第%d课' % order_no)


def _load_import_meta(token):
    if not token or not re.match(r'^[a-f0-9]{16}$', token):
        return None
    path = os.path.join(TMP_DIR, token + '.json')
    if not os.path.exists(path):
        return None
    with open(path, 'r', encoding='utf-8') as fp:
        return json.load(fp)


def _cleanup_import(token):
    for suffix in ('.json', '.xlsx', '.doc', '.docx'):
        p = os.path.join(TMP_DIR, token + suffix)
        if os.path.exists(p):
            try:
                os.remove(p)
            except Exception:
                pass


@lesson_prep_bp.route('/periods/template')
@login_required
def periods_template():
    return send_file(_build_progress_template(), as_attachment=True,
                     download_name='教学进度表导入模板.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@lesson_prep_bp.route('/periods/import/parse', methods=['POST'])
@login_required
def periods_import_parse():
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'ok': False, 'error': '请选择要上传的文件'}), 400
    ext = os.path.splitext(f.filename)[1].lower()
    if ext not in ALLOWED_PROGRESS_EXT:
        return jsonify({'ok': False, 'error': '仅支持 xlsx / doc / docx 格式文件'}), 400
    data = f.read()
    if len(data) > 20 * 1024 * 1024:
        return jsonify({'ok': False, 'error': '文件大小不能超过 20MB'}), 400

    if ext == '.xlsx':
        rows, error, warning = _parse_progress_xlsx(io.BytesIO(data))
    else:
        rows, error, warning = _parse_progress_docx(io.BytesIO(data))
    if error:
        if ext == '.doc':
            error = '检测到旧版 .doc 文档，自动解析失败，请另存为 .docx 或下载标准模板重新上传'
        return jsonify({'ok': False, 'error': error}), 400

    token = uuid.uuid4().hex[:16]
    with open(os.path.join(TMP_DIR, token + ext), 'wb') as fp:
        fp.write(data)
    meta = {'file_name': os.path.basename(f.filename), 'ext': ext, 'warning': warning, 'rows': rows}
    with open(os.path.join(TMP_DIR, token + '.json'), 'w', encoding='utf-8') as fp:
        json.dump(meta, fp, ensure_ascii=False)
    return jsonify({'ok': True, 'token': token, 'rows': rows, 'warning': warning,
                    'file_name': meta['file_name']})


@lesson_prep_bp.route('/periods/class-existing')
@login_required
def periods_class_existing():
    """返回指定班级已导入（schedule_import）课次编号列表，供预览页追加模式冲突校验。"""
    class_id = request.args.get('class_id', type=int)
    orders = []
    if class_id:
        orders = sorted(p.order_no for p in ClassPeriod.query.filter_by(
            class_id=class_id, source_type='schedule_import').all() if p.order_no is not None)
    return jsonify({'orders': orders})


@lesson_prep_bp.route('/periods/import/confirm', methods=['POST'])
@login_required
def periods_import_confirm():
    data = request.get_json(silent=True)
    if data is None:
        data = request.form
    token = data.get('token', '')
    meta = _load_import_meta(token)
    if not meta:
        return jsonify({'ok': False, 'error': '导入会话已失效，请重新上传'}), 400
    # 班级id数组（pclass_ids），一份进度表可同步至多个班级
    pclass_ids = data.get('pclass_ids', [])
    if isinstance(pclass_ids, (str, int)):
        pclass_ids = [pclass_ids]
    pclass_ids = [int(c) for c in pclass_ids if str(c).strip()]
    course_name = (data.get('course_name') or '').strip()
    sport_item = (data.get('sport_item') or '').strip()
    location = (data.get('location') or '').strip()
    mode = data.get('mode', 'append')  # append / cover
    rows = data.get('rows', [])
    if isinstance(rows, str):
        try:
            rows = json.loads(rows)
        except Exception:
            rows = []
    if not pclass_ids:
        return jsonify({'ok': False, 'error': '请选择绑定班级'}), 400
    if not course_name:
        return jsonify({'ok': False, 'error': '请选择课程名称'}), 400
    if not sport_item:
        return jsonify({'ok': False, 'error': '请选择授课项目'}), 400
    if mode not in ('append', 'cover'):
        mode = 'append'
    if not rows:
        return jsonify({'ok': False, 'error': '没有可导入的课次数据'}), 400

    # 源文件最终路径（提交成功后再移动，避免事务回滚留下孤儿文件）
    final_name = 'period_%s%s' % (uuid.uuid4().hex[:10], meta['ext'])
    dest = os.path.join(UPLOAD_DIR, final_name)
    tmp_file = os.path.join(TMP_DIR, token + meta['ext'])
    source_file_id = None

    imported_classes = 0
    imported_periods = 0
    try:
        if os.path.exists(tmp_file):
            r = Resource(name=meta['file_name'], category='教学进度表',
                         file_path='/static/uploads/' + final_name,
                         note='进度表导入（%s）' % course_name)
            db.session.add(r)
            db.session.flush()
            source_file_id = r.id
        for cid in pclass_ids:
            # 冲突校验：该班已存在 schedule_import 来源的课次（手动新增课时不参与冲突）
            existing = set(p.order_no for p in ClassPeriod.query.filter_by(
                class_id=cid, source_type='schedule_import').all() if p.order_no is not None)
            seen = set()
            for row in rows:
                try:
                    o = int(row.get('order_no'))
                except (TypeError, ValueError):
                    raise ValueError('课次必须为数字')
                if o in seen:
                    raise ValueError('课次 %d 在本次导入中重复' % o)
                seen.add(o)
                if mode == 'append' and o in existing:
                    raise ValueError('课次 %d 与该班级现有进度表课时冲突' % o)
            # cover：先删除该班进度表导入课时（手动新增课时保留）
            if mode == 'cover':
                ClassPeriod.query.filter_by(class_id=cid, source_type='schedule_import').delete()
            for row in rows:
                try:
                    o = int(row.get('order_no'))
                except (TypeError, ValueError):
                    continue
                content = (row.get('content') or '').strip()
                if not content:
                    continue
                sport = (row.get('sport_item') or sport_item).strip()
                loc = (row.get('location') or '').strip()
                if not loc:
                    loc = SPORT_VENUES.get(sport, '')
                db.session.add(ClassPeriod(
                    class_id=cid, order_no=o, name=_period_label(o, content),
                    content=content, week=(row.get('week') or '').strip(),
                    section=(row.get('section') or '').strip(), hours=(row.get('hours') or '').strip(),
                    homework=(row.get('homework') or '').strip(), course_name=course_name,
                    sport_item=sport, location=loc, note=(row.get('note') or '').strip(),
                    source_file_id=source_file_id, source_type='schedule_import',
                ))
                imported_periods += 1
            imported_classes += 1
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        return jsonify({'ok': False, 'error': '导入失败，已整体回滚：%s' % e}), 500

    # 提交成功后移动源文件到资源库目录
    if os.path.exists(tmp_file):
        try:
            os.replace(tmp_file, dest)
        except Exception:
            pass
    _cleanup_import(token)
    return jsonify({'ok': True, 'imported_classes': imported_classes,
                    'imported_periods': imported_periods, 'first_class_id': pclass_ids[0]})


@lesson_prep_bp.route('/periods/import/cancel', methods=['POST'])
@login_required
def periods_import_cancel():
    data = request.get_json(silent=True) or request.form
    token = data.get('token', '')
    _cleanup_import(token)
    return jsonify({'ok': True})


SPORT_LESSON_TEMPLATE = """【课题】

一、教学目标
1. 知识与技能：
2. 过程与方法：
3. 情感态度与价值观：

二、教学重点与难点
重点：
难点：

三、教学准备
场地：
器材：

四、教学过程
（一）准备部分（约  分钟）
（二）基本部分（约  分钟）
（三）结束部分（约  分钟）

五、教学评价与安全注意事项
"""


def _all_classes():
    return SchoolClass.query.order_by(SchoolClass.id).all()


@lesson_prep_bp.route('/plans')
@login_required
def plans():
    from sqlalchemy import func
    keyword = request.args.get('keyword', '').strip()
    category = request.args.get('category', '')
    sport_item = request.args.get('sport_item', '')
    grade = request.args.get('grade', '')
    sub = db.session.query(LessonPlan.group_id, func.max(LessonPlan.id).label('max_id')).group_by(LessonPlan.group_id).subquery()
    q = db.session.query(LessonPlan).join(sub, LessonPlan.id == sub.c.max_id)
    if keyword:
        q = q.filter(db.or_(LessonPlan.title.like('%%%s%%' % keyword),
                            LessonPlan.subject.like('%%%s%%' % keyword),
                            LessonPlan.content.like('%%%s%%' % keyword)))
    if category:
        q = q.filter(LessonPlan.category == category)
    if sport_item:
        q = q.filter(LessonPlan.sport_item == sport_item)
    if grade:
        q = q.filter(LessonPlan.grade == grade)
    plans_list = q.order_by(LessonPlan.id.desc()).all()
    return render_template('lesson_prep/plans.html', active='plans', plans=plans_list,
                           categories=RESOURCE_CATEGORIES, sport_items=SPORT_ITEMS, grades=GRADES,
                           keyword=keyword, cur_category=category, cur_sport=sport_item, cur_grade=grade)


@lesson_prep_bp.route('/plans/new', methods=['GET', 'POST'])
@login_required
def plan_new():
    if request.method == 'POST':
        title = request.form.get('title', '').strip()
        content = request.form.get('content', '')
        if not title:
            flash('请填写教案标题', 'error')
            return redirect(url_for('lesson_prep.plan_new'))
        db.session.add(LessonPlan(
            group_id=uuid.uuid4().hex[:12], version=1, title=title,
            subject=request.form.get('subject', '').strip(),
            class_id=request.form.get('class_id', type=int),
            duration=request.form.get('duration', '').strip(),
            category=request.form.get('category', '教案模板'),
            sport_item=request.form.get('sport_item', '其他项目'),
            grade=request.form.get('grade', '通用'),
            content=content,
        ))
        db.session.commit()
        flash('教案已保存', 'success')
        return redirect(url_for('lesson_prep.plans'))
    return render_template('lesson_prep/plan_edit.html', active='plans', plan=None,
                           classes=_all_classes(), template=SPORT_LESSON_TEMPLATE,
                           categories=RESOURCE_CATEGORIES, sport_items=SPORT_ITEMS, grades=GRADES)


@lesson_prep_bp.route('/plans/<int:pid>/edit', methods=['GET', 'POST'])
@login_required
def plan_edit(pid):
    plan = db.get_or_404(LessonPlan, pid)
    latest = LessonPlan.query.filter_by(group_id=plan.group_id).order_by(LessonPlan.version.desc()).first()
    if request.method == 'POST':
        new_plan = LessonPlan(
            group_id=plan.group_id,
            version=(latest.version + 1) if latest else 1,
            title=request.form.get('title', '').strip(),
            subject=request.form.get('subject', '').strip(),
            class_id=request.form.get('class_id', type=int),
            duration=request.form.get('duration', '').strip(),
            category=request.form.get('category', plan.category or '教案模板'),
            sport_item=request.form.get('sport_item', plan.sport_item or '其他项目'),
            grade=request.form.get('grade', plan.grade or '通用'),
            content=request.form.get('content', ''),
        )
        db.session.add(new_plan)
        db.session.commit()
        flash('已保存新版本 v%d' % new_plan.version, 'success')
        return redirect(url_for('lesson_prep.plans'))
    quality = request.args.get('quality', '')
    original_url = request.args.get('original_url', '')
    return render_template('lesson_prep/plan_edit.html', active='plans', plan=latest or plan,
                           classes=_all_classes(), template=SPORT_LESSON_TEMPLATE,
                           categories=RESOURCE_CATEGORIES, sport_items=SPORT_ITEMS, grades=GRADES,
                           quality=quality, original_url=original_url)


@lesson_prep_bp.route('/plans/<int:pid>/view')
@login_required
def plan_view(pid):
    plan = db.get_or_404(LessonPlan, pid)
    return render_template('lesson_prep/plan_view.html', active='plans', plan=plan)


@lesson_prep_bp.route('/plans/<int:pid>/history')
@login_required
def plan_history(pid):
    plan = db.get_or_404(LessonPlan, pid)
    versions = LessonPlan.query.filter_by(group_id=plan.group_id).order_by(LessonPlan.version.desc()).all()
    return render_template('lesson_prep/plan_history.html', active='plans', plan=plan, versions=versions)


@lesson_prep_bp.route('/plans/<int:pid>/rollback/<int:vid>', methods=['POST'])
@login_required
def plan_rollback(pid, vid):
    plan = db.get_or_404(LessonPlan, pid)
    target = db.get_or_404(LessonPlan, vid)
    latest = LessonPlan.query.filter_by(group_id=plan.group_id).order_by(LessonPlan.version.desc()).first()
    new_plan = LessonPlan(group_id=plan.group_id, version=(latest.version + 1) if latest else 1,
                          title=target.title, subject=target.subject, class_id=target.class_id,
                          duration=target.duration, content=target.content)
    db.session.add(new_plan)
    db.session.commit()
    flash('已还原到 v%d（生成新版本 v%d）' % (target.version, new_plan.version), 'success')
    return redirect(url_for('lesson_prep.plan_history', pid=plan.id))


@lesson_prep_bp.route('/plans/compare/<group_id>')
@login_required
def plan_compare(group_id):
    versions = LessonPlan.query.filter_by(group_id=group_id).order_by(LessonPlan.version).all()
    v1_id = request.args.get('v1', type=int)
    v2_id = request.args.get('v2', type=int)
    v1 = v2 = None
    if versions:
        v2 = next((v for v in versions if v.id == v2_id), versions[-1])
        v1 = next((v for v in versions if v.id == v1_id), versions[-2] if len(versions) >= 2 else versions[-1])
    diff = ''
    if v1 and v2:
        diff = difflib.HtmlDiff(wrapcolumn=60).make_table(
            (v1.content or '').splitlines(), (v2.content or '').splitlines(),
            fromdesc='v%d' % v1.version, todesc='v%d' % v2.version)
    return render_template('lesson_prep/plan_compare.html', active='plans',
                           versions=versions, v1=v1, v2=v2, diff=diff)


@lesson_prep_bp.route('/plans/<int:pid>/delete', methods=['POST'])
@login_required
def plan_delete(pid):
    plan = db.get_or_404(LessonPlan, pid)
    LessonPlan.query.filter_by(group_id=plan.group_id).delete()
    db.session.commit()
    flash('教案已删除（含全部版本）', 'success')
    return redirect(url_for('lesson_prep.plans'))


def _docx_to_html(data):
    """用 mammoth 把 docx 转成 HTML 富文本（严格保留原位表格、图片）。返回 (html, warnings)。"""
    import io
    import mammoth

    def _img_cb(image):
        with image.open() as f:
            blob = f.read()
        ext = '.png'
        if image.content_type:
            m = image.content_type.split('/')[-1].lower()
            if m in ('png', 'jpg', 'jpeg', 'gif', 'bmp', 'webp'):
                ext = '.' + m
        filename = 'planimg_%s%s' % (uuid.uuid4().hex[:10], ext)
        with open(os.path.join(UPLOAD_DIR, filename), 'wb') as fp:
            fp.write(blob)
        return {'src': '/static/uploads/' + filename}

    result = mammoth.convert_to_html(
        io.BytesIO(data),
        convert_image=mammoth.images.img_element(_img_cb),
    )
    html = result.value
    warnings = [m.message for m in result.messages]
    # 调试打印：确认表格/图片位置
    tbl_pos = html.find('<table')
    si_pos = html.find('四、教学设计与过程')
    img_pos = html.find('<img')
    print('[docx解析] HTML长度=%d 表格位置=%d 四、教学设计与过程位置=%d 图片位置=%d' % (
        len(html), tbl_pos, si_pos, img_pos), flush=True)
    if si_pos != -1 and tbl_pos != -1:
        print('[docx解析] 表格在「四、教学设计与过程」之后: %s' % ('是' if tbl_pos > si_pos else '否'), flush=True)
    return html, warnings


def _parse_plan_file(data, ext):
    """解析上传的教案文件，返回 (HTML, quality, warnings)；失败返回 (None, 'none', [])。"""
    if ext == '.docx':
        try:
            html, warnings = _docx_to_html(data)
            if not html or not html.strip():
                return None, 'none', []
            quality = 'partial' if warnings else 'full'
            return html, quality, warnings
        except Exception:
            return None, 'none', []
    if ext == '.pdf':
        # PDF 仅预览：不提取文字，保存原文件并提供预览链接
        filename = 'planpdf_%s.pdf' % uuid.uuid4().hex[:10]
        with open(os.path.join(UPLOAD_DIR, filename), 'wb') as fp:
            fp.write(data)
        return '<p>（PDF 文档，仅预览，不支持在线编辑）</p><p><a href="/static/uploads/' + filename + '" target="_blank" class="btn btn-outline">打开 PDF 原版预览</a></p>', 'full', []
    if ext in ('.png', '.jpg', '.jpeg'):
        filename = 'planimg_%s%s' % (uuid.uuid4().hex[:10], ext)
        with open(os.path.join(UPLOAD_DIR, filename), 'wb') as fp:
            fp.write(data)
        return '<p><img src="/static/uploads/' + filename + '" style="max-width:100%"></p><p>（图片已嵌入教案正文，可在此下方补充文字说明）</p>', 'full', []
    return None, 'none', []


@lesson_prep_bp.route('/plans/import', methods=['POST'])
@login_required
def plan_import():
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'error': '请选择文件'}), 400
    ext = os.path.splitext(f.filename)[1].lower()
    if ext not in ('.docx', '.pdf', '.png', '.jpg', '.jpeg'):
        return jsonify({'error': '仅支持docx、pdf、png、jpg、jpeg格式文件'}), 400
    data = f.read()
    if len(data) > 20 * 1024 * 1024:
        return jsonify({'error': '文件大小不能超过20MB'}), 400
    content, quality, warnings = _parse_plan_file(data, ext)
    if content is None:
        return jsonify({'error': '文档解析失败，请检查文件内容后重试'}), 400
    name = request.form.get('name', '').strip() or os.path.splitext(os.path.basename(f.filename))[0]
    category = request.form.get('category', '教案模板')
    sport_item = request.form.get('sport_item', '其他项目')
    grade = request.form.get('grade', '通用')
    original_url = ''
    if ext == '.docx' and quality == 'partial':
        orig_name = 'planorig_%s.docx' % uuid.uuid4().hex[:10]
        with open(os.path.join(UPLOAD_DIR, orig_name), 'wb') as fp:
            fp.write(data)
        original_url = '/static/uploads/' + orig_name
    plan = LessonPlan(
        group_id=uuid.uuid4().hex[:12], version=1, title=name,
        subject=sport_item, category=category, sport_item=sport_item, grade=grade,
        content=content,
    )
    db.session.add(plan)
    db.session.commit()
    return jsonify({'ok': True, 'plan_id': plan.id, 'content': content,
                    'parse_quality': quality, 'warn_messages': warnings,
                    'original_url': original_url,
                    'message': '教案导入成功，已自动创建初始版本并绑定分类信息'})


# ============ 学情分层作业中心 ============
LEVELS = ['A层（提高层）', 'B层（基础层）', 'C层（巩固层）']


def _compute_level(fitness, skill):
    """总分 = 体能分*0.5 + 技能分*0.5；A层≥80，B层60~80，C层<60。"""
    total = (fitness + skill) / 2.0
    if total >= 80:
        return 'A层（提高层）'
    if total >= 60:
        return 'B层（基础层）'
    return 'C层（巩固层）'


@lesson_prep_bp.route('/layered')
@login_required
def layered():
    classes = _all_classes()
    cid = request.args.get('class_id', type=int)
    if cid is None and classes:
        cid = classes[0].id
    selected = db.session.get(SchoolClass, cid) if cid else None
    # 运动项目下拉数据源：取自「班级-运动项目关联配置」，不允许在分层页直接新增绑定
    sport_items = []
    if cid:
        sport_items = [b.sport_item for b in ClassSportBinding.query
                       .filter_by(class_id=cid).order_by(ClassSportBinding.id).all()]
    sport_item = request.args.get('sport_item', '').strip()
    if sport_item not in sport_items:
        sport_item = sport_items[0] if sport_items else ''
    # 关联课时下拉数据源：取自「课时管理」内该行政班的课时
    periods = []
    if cid:
        periods = ClassPeriod.query.filter_by(class_id=cid).order_by(ClassPeriod.order_no, ClassPeriod.id).all()
    period_id = request.args.get('period_id', type=int)
    if period_id and not any(p.id == period_id for p in periods):
        period_id = None
    if period_id is None and periods:
        period_id = periods[0].id
    students = Student.query.filter_by(class_id=cid).order_by(Student.student_no).all() if cid else []
    # 学情按 行政班 + 运动项目 + 课时 三元组隔离
    learning_map = {}
    if students and period_id and sport_item:
        sids = [s.id for s in students]
        for sl in StudentLearning.query.filter(
                StudentLearning.class_id == cid,
                StudentLearning.sport_item == sport_item,
                StudentLearning.period_id == period_id,
                StudentLearning.student_id.in_(sids)).all():
            learning_map[sl.student_id] = sl
    return render_template('lesson_prep/layered.html', active='layered',
                           classes=classes, selected=selected, students=students,
                           sport_items=sport_items, sport_item=sport_item,
                           periods=periods, period_id=period_id,
                           learning_map=learning_map, levels=LEVELS)


@lesson_prep_bp.route('/layered/save', methods=['POST'])
@login_required
def layered_save():
    cid = request.form.get('class_id', type=int)
    sport_item = request.form.get('sport_item', '').strip()
    period_id = request.form.get('period_id', type=int)
    if not sport_item:
        flash('该班级尚未配置运动项目，请先到「教师工作台 → 学生名册 → 班级管理」绑定后再做学情分层', 'error')
        return redirect(url_for('lesson_prep.layered', class_id=cid, period_id=period_id))
    if not period_id:
        flash('该班级尚未配置课时，请先到「教师工作台 → 教学进度」添加课时后再做学情分层', 'error')
        return redirect(url_for('lesson_prep.layered', class_id=cid, sport_item=sport_item))
    students = Student.query.filter_by(class_id=cid).all()
    for s in students:
        fitness = request.form.get('fitness_%d' % s.id, type=float)
        skill = request.form.get('skill_%d' % s.id, type=float)
        level_manual = request.form.get('level_%d' % s.id, '').strip()
        if fitness is None and skill is None and not level_manual:
            continue
        fitness = fitness if fitness is not None else 0.0
        skill = skill if skill is not None else 0.0
        # 允许教师手动覆盖分层，未选择时按规则自动计算
        level = level_manual if level_manual in LEVELS else _compute_level(fitness, skill)
        rec = StudentLearning.query.filter_by(class_id=cid, sport_item=sport_item,
                                              period_id=period_id, student_id=s.id).first()
        if rec:
            rec.fitness_score = fitness
            rec.skill_score = skill
            rec.level = level
        else:
            db.session.add(StudentLearning(class_id=cid, sport_item=sport_item, period_id=period_id,
                                           student_id=s.id, fitness_score=fitness, skill_score=skill, level=level))
    db.session.commit()
    flash('学情数据已保存并自动分层', 'success')
    return redirect(url_for('lesson_prep.layered', class_id=cid, sport_item=sport_item, period_id=period_id))


@lesson_prep_bp.route('/layered/batch')
@login_required
def layered_batch():
    """批量分层：按 行政班+运动项目+课时 筛选，批量计算所有学生层次并返回。"""
    cid = request.args.get('class_id', type=int)
    sport_item = request.args.get('sport_item', '').strip()
    period_id = request.args.get('period_id', type=int)
    items = []
    if cid and sport_item and period_id:
        students = Student.query.filter_by(class_id=cid).order_by(Student.student_no).all()
        for s in students:
            sl = StudentLearning.query.filter_by(class_id=cid, sport_item=sport_item,
                                                 period_id=period_id, student_id=s.id).first()
            fitness = sl.fitness_score if sl else None
            skill = sl.skill_score if sl else None
            if fitness is None or skill is None:
                continue
            items.append({'student_id': s.id, 'student_no': s.student_no, 'name': s.name,
                          'fitness': fitness, 'skill': skill,
                          'total': round((fitness + skill) / 2.0, 1),
                          'level': _compute_level(fitness, skill)})
    return jsonify({'items': items})


@lesson_prep_bp.route('/layered/assignments/save', methods=['POST'])
@login_required
def layered_assignments_save():
    period_id = request.form.get('period_id', type=int)
    sport_item = request.form.get('sport_item', '').strip()
    level = request.form.get('level', LEVELS[0])
    content = request.form.get('content', '').strip()
    if content:
        db.session.add(LayeredAssignment(period_id=period_id, sport_item=sport_item,
                                         level=level, content=content))
        db.session.commit()
        flash('分层作业已保存', 'success')
    return redirect(url_for('lesson_prep.layered', class_id=request.form.get('class_id', type=int),
                            sport_item=sport_item, period_id=period_id))


# ============ 备课成果管理 ============

@lesson_prep_bp.route('/outcomes')
@login_required
def outcomes():
    from sqlalchemy import func
    keyword = request.args.get('keyword', '').strip()
    sport_item = request.args.get('sport_item', '')
    grade = request.args.get('grade', '')
    sub = db.session.query(LessonPlan.group_id, func.max(LessonPlan.id).label('max_id')).group_by(LessonPlan.group_id).subquery()
    q = db.session.query(LessonPlan).join(sub, LessonPlan.id == sub.c.max_id)
    if keyword:
        q = q.filter(db.or_(LessonPlan.title.like('%%%s%%' % keyword),
                            LessonPlan.subject.like('%%%s%%' % keyword),
                            LessonPlan.content.like('%%%s%%' % keyword)))
    if sport_item:
        q = q.filter(LessonPlan.sport_item == sport_item)
    if grade:
        q = q.filter(LessonPlan.grade == grade)
    plans_list = q.order_by(LessonPlan.id.desc()).all()
    plan_ids = [p.id for p in plans_list]
    sheet_counts = {}
    if plan_ids:
        for row in db.session.query(LearningSheet.lesson_plan_id, func.count(LearningSheet.id)).filter(
                LearningSheet.lesson_plan_id.in_(plan_ids)).group_by(LearningSheet.lesson_plan_id).all():
            sheet_counts[row[0]] = row[1]
    version_counts = {}
    if plans_list:
        gids = [p.group_id for p in plans_list]
        for row in db.session.query(LessonPlan.group_id, func.count(LessonPlan.id)).filter(
                LessonPlan.group_id.in_(gids)).group_by(LessonPlan.group_id).all():
            version_counts[row[0]] = row[1]
    return render_template('lesson_prep/outcomes.html', active='outcomes',
                           plans=plans_list, sheet_counts=sheet_counts, version_counts=version_counts,
                           sport_items=SPORT_ITEMS, grades=GRADES,
                           keyword=keyword, cur_sport=sport_item, cur_grade=grade)


@lesson_prep_bp.route('/outcomes/<int:pid>/sheets')
@login_required
def outcome_sheets(pid):
    plan = db.get_or_404(LessonPlan, pid)
    sheets = LearningSheet.query.filter_by(lesson_plan_id=pid).order_by(LearningSheet.id).all()
    sheets_map = {s.phase: s for s in sheets}
    return render_template('lesson_prep/outcome_sheets.html', active='outcomes',
                           plan=plan, sheets_map=sheets_map, providers=provider_options())


@lesson_prep_bp.route('/outcomes/<int:pid>/sheets/save', methods=['POST'])
@login_required
def outcome_sheets_save(pid):
    for phase in ('课前', '课中', '课后'):
        content = request.form.get('content_%s' % phase, '').strip()
        rec = LearningSheet.query.filter_by(lesson_plan_id=pid, phase=phase).first()
        if content:
            if rec:
                rec.content = content
            else:
                db.session.add(LearningSheet(lesson_plan_id=pid, phase=phase, content=content))
        elif rec:
            db.session.delete(rec)
    db.session.commit()
    flash('学习单已保存', 'success')
    return redirect(url_for('lesson_prep.outcome_sheets', pid=pid))


@lesson_prep_bp.route('/outcomes/<int:pid>/sheets/export')
@login_required
def outcome_sheets_export(pid):
    plan = db.get_or_404(LessonPlan, pid)
    sheets = LearningSheet.query.filter_by(lesson_plan_id=pid).all()
    order = {'课前': 0, '课中': 1, '课后': 2}
    sheets = sorted(sheets, key=lambda s: order.get(s.phase, 9))
    lines = ['# %s 一体化学习单' % plan.title, '主题：%s' % (plan.subject or ''), '']
    for s in sheets:
        lines.append('## %s学习单' % s.phase)
        lines.append(s.content or '')
        lines.append('')
    buf = io.BytesIO(('\n'.join(lines)).encode('utf-8'))
    return send_file(buf, as_attachment=True, download_name='%s-学习单.txt' % plan.title,
                     mimetype='text/plain')

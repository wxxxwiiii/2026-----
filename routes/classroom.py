"""课堂智慧助手：学生考勤+随机抽提问、学生平时成绩（考勤+答题，满分100）。"""
import io
import random
from collections import Counter
from datetime import date, datetime

from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify, send_file

from extensions import db
from models import SchoolClass, Student, Attendance, AnswerRecord, GradeSetting, RollCallRecord
from security import login_required
from services.settings import get_setting, set_setting

classroom_bp = Blueprint('classroom', __name__, url_prefix='/classroom')

ATT_WEIGHTS = {'出勤': 1.0, '迟到': 0.6, '请假': 0.4, '缺勤': 0.0}
ANSWER_WEIGHTS = {'答对': 1.0, '部分': 0.5, '答错': 0.0, '未答': 0.0}


def _classes():
    return SchoolClass.query.order_by(SchoolClass.id).all()


def _get_weights():
    w = {s.key: s.weight for s in GradeSetting.query.all()}
    return {'attendance': w.get('attendance', 60), 'answer': w.get('answer', 40)}


def _grade_label(total):
    if total >= 90:
        return '优秀'
    if total >= 80:
        return '良好'
    if total >= 60:
        return '合格'
    return '不及格'


def _year_options():
    current = date.today().year
    return [str(y) for y in range(current, current - 10, -1)]


def compute_grades(class_id, weights, year=None):
    """计算某班平时成绩（可按学年过滤考勤/作答记录）。"""
    students = Student.query.filter_by(class_id=class_id).order_by(Student.student_no).all()
    grades = []
    for s in students:
        att_q = Attendance.query.filter_by(student_id=s.id)
        ans_q = AnswerRecord.query.filter_by(student_id=s.id)
        if year:
            att_q = att_q.filter(Attendance.date.like(year + '%'))
            ans_q = ans_q.filter(AnswerRecord.created_at >= datetime(int(year), 1, 1),
                                 AnswerRecord.created_at < datetime(int(year) + 1, 1, 1))
        atts = att_q.all()
        att_avg = (sum(ATT_WEIGHTS.get(a.status, 0.0) for a in atts) / len(atts)) if atts else 0.0
        answers = ans_q.all()
        ans_avg = (sum(ANSWER_WEIGHTS.get(a.result, 0.0) for a in answers) / len(answers)) if answers else 0.0
        att_score = weights['attendance'] * att_avg
        ans_score = weights['answer'] * ans_avg
        total = round(att_score + ans_score, 1)
        grades.append({
            'student_no': s.student_no,
            'name': s.name,
            'attendance_score': round(att_score, 1),
            'answer_score': round(ans_score, 1),
            'total': total,
            'rating': _grade_label(total),
        })
    return grades


@classroom_bp.route('/')
@login_required
def index():
    return redirect(url_for('classroom.attendance'))


@classroom_bp.route('/attendance')
@login_required
def attendance():
    """学生考勤 + 随机抽提问页面。"""
    classes = _classes()
    cid = request.args.get('class_id', type=int)
    if cid is None and classes:
        cid = classes[0].id
    selected = db.session.get(SchoolClass, cid) if cid else None
    students = Student.query.filter_by(class_id=cid).order_by(Student.student_no).all() if cid else []
    att_date = request.args.get('date') or date.today().isoformat()

    student_ids = {s.id for s in students}
    attendance_map = {}
    if student_ids:
        recs = Attendance.query.filter(Attendance.date == att_date,
                                       Attendance.student_id.in_(student_ids)).all()
        attendance_map = {r.student_id: r.status for r in recs}
    stats = Counter(attendance_map.values())

    return render_template('classroom/attendance.html', active='classroom_attendance',
                           classes=classes, selected=selected, students=students,
                           att_date=att_date, attendance_map=attendance_map, stats=stats)


@classroom_bp.route('/grades')
@login_required
def grades():
    """学生平时成绩页面（班级 + 学年筛选）。"""
    classes = _classes()
    cid = request.args.get('class_id', type=int)
    if cid is None and classes:
        cid = classes[0].id
    year = request.args.get('year', '').strip()
    selected = db.session.get(SchoolClass, cid) if cid else None
    weights = _get_weights()
    grades = compute_grades(cid, weights, year=year) if cid else []
    return render_template('classroom/grades.html', active='classroom_grades',
                           classes=classes, selected=selected, year=year,
                           weights=weights, grades=grades, year_options=_year_options())


@classroom_bp.route('/attendance', methods=['POST'])
@login_required
def attendance_save():
    cid = request.form.get('class_id', type=int)
    att_date = request.form.get('date') or date.today().isoformat()
    students = Student.query.filter_by(class_id=cid).all()
    for s in students:
        status = request.form.get('status_%d' % s.id, '出勤')
        rec = Attendance.query.filter_by(student_id=s.id, date=att_date).first()
        if rec:
            rec.status = status
        else:
            db.session.add(Attendance(student_id=s.id, date=att_date, status=status))
    db.session.commit()
    flash('考勤已保存', 'success')
    return redirect(url_for('classroom.attendance', class_id=cid, date=att_date))


def _current_round(class_id):
    """当前抽取轮次（按班级隔离，默认第 1 轮）。"""
    try:
        return int(get_setting('pick_round_%d' % class_id, '1') or '1')
    except (TypeError, ValueError):
        return 1


@classroom_bp.route('/api/pick')
@login_required
def api_pick():
    """随机抽取学生：本轮已抽过的学生不重复抽取。"""
    cid = request.args.get('class_id', type=int)
    students = Student.query.filter_by(class_id=cid).order_by(Student.student_no).all()
    if not students:
        return jsonify({'error': '该班级没有学生，请先到教师工作台录入学生'})
    rnd = _current_round(cid)
    picked_ids = {r.student_id for r in RollCallRecord.query.filter_by(class_id=cid, round_no=rnd).all()}
    candidates = [s for s in students if s.id not in picked_ids]
    if not candidates:
        return jsonify({'error': '本轮学生均已抽取，请点击「重置抽取记录」开始新一轮'})
    s = random.choice(candidates)
    db.session.add(RollCallRecord(class_id=cid, student_id=s.id, mode='追问', round_no=rnd))
    db.session.commit()
    picked_ids.add(s.id)
    return jsonify({'id': s.id, 'name': s.name, 'student_no': s.student_no or '',
                    'round': rnd, 'picked_ids': sorted(picked_ids)})


@classroom_bp.route('/api/pick/reset', methods=['POST'])
@login_required
def api_pick_reset():
    """重置抽取记录：开启新一轮，清空本轮已抽取名单。"""
    cid = request.form.get('class_id', type=int)
    rnd = _current_round(cid) + 1
    set_setting('pick_round_%d' % cid, str(rnd))
    return jsonify({'round': rnd, 'picked_ids': []})


@classroom_bp.route('/api/pick/records')
@login_required
def api_pick_records():
    """读取当前轮已抽取名单 + 历史抽问记录。"""
    cid = request.args.get('class_id', type=int)
    rnd = _current_round(cid)
    records = RollCallRecord.query.filter_by(class_id=cid).order_by(RollCallRecord.id.desc()).all()
    picked_ids = [r.student_id for r in records if (r.round_no or 1) == rnd]
    history = []
    for r in records[:100]:
        s = db.session.get(Student, r.student_id)
        history.append({'student_id': r.student_id, 'name': s.name if s else '未知',
                        'student_no': s.student_no if s else '', 'round': r.round_no or 1,
                        'time': r.created_at.strftime('%Y-%m-%d %H:%M:%S') if r.created_at else ''})
    return jsonify({'round': rnd, 'picked_ids': picked_ids, 'history': history})


@classroom_bp.route('/answer', methods=['POST'])
@login_required
def answer_save():
    cid = request.form.get('class_id', type=int)
    sid = request.form.get('student_id', type=int)
    result = request.form.get('result', '未答')
    if sid:
        db.session.add(AnswerRecord(class_id=cid, student_id=sid, result=result))
        db.session.commit()
        flash('已记录作答情况', 'success')
    return redirect(url_for('classroom.attendance', class_id=cid))


@classroom_bp.route('/weights', methods=['POST'])
@login_required
def weights_save():
    cid = request.form.get('class_id', type=int)
    try:
        att = float(request.form.get('attendance_weight', 60))
        ans = float(request.form.get('answer_weight', 40))
    except ValueError:
        att, ans = 60, 40
    for key, val, name in (('attendance', att, '考勤'), ('answer', ans, '课堂答题')):
        gs = GradeSetting.query.filter_by(key=key).first()
        if gs:
            gs.weight = val
        else:
            db.session.add(GradeSetting(key=key, name=name, weight=val))
    db.session.commit()
    flash('平时成绩权重已保存', 'success')
    return redirect(url_for('classroom.grades', class_id=cid))


@classroom_bp.route('/grades/export')
@login_required
def grades_export():
    """导出当前筛选后的班级平时成绩 Excel。"""
    cid = request.args.get('class_id', type=int)
    year = request.args.get('year', '').strip()
    cls = db.session.get(SchoolClass, cid) if cid else None
    weights = _get_weights()
    grades = compute_grades(cid, weights, year=year) if cid else []
    buf = _build_grades_excel(cls, grades, year)
    return send_file(buf, as_attachment=True,
                     download_name='%s-平时成绩%s.xlsx' % (cls.name if cls else '班级', ('-' + year) if year else ''),
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def _build_grades_excel(cls, grades, year):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = '平时成绩'
    ws.cell(row=1, column=1, value='%s 平时成绩（%s）' % (
        cls.name if cls else '', year or '全部学年')).font = Font(bold=True, size=14)
    headers = ['学号', '姓名', '考勤分', '答题分', '总分', '评级']
    for i, h in enumerate(headers, start=1):
        ws.cell(row=3, column=i, value=h).font = Font(bold=True)
    row = 4
    for g in grades:
        vals = [g['student_no'] or '', g['name'], g['attendance_score'],
                g['answer_score'], g['total'], g['rating']]
        for i, v in enumerate(vals, start=1):
            ws.cell(row=row, column=i, value=v)
        row += 1
    for i, w in enumerate((14, 12, 10, 10, 10, 10), start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

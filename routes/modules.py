"""体质健康分析（5 个子页面）、AI 动作识别（6 个子菜单）。"""
import csv
import datetime
import io
import json
import os
import re
import threading
import uuid

from flask import Blueprint, render_template, request, redirect, url_for, flash, send_file, jsonify, current_app

from config import UPLOAD_DIR
from extensions import db
from models import (SchoolClass, Student, FitnessRecord, ActionAnalysis, TrainingPlan,
                    SPORT_ITEMS, SPORT_TECHNIQUES, ACTION_STATUS_LABELS)
from security import login_required
from services.fitness_scoring import calculate
from services.vision import extract_frames, analyze_action
from services.pose_analysis import analyze_pose, MP_AVAILABLE
from services.aggregation import summarize_class_errors, AGGREGATION_MODEL
from services.ai import get_client

modules_bp = Blueprint('modules', __name__)

ITEM_LABELS = [
    ('bmi', '体重指数（BMI）'),
    ('vital_capacity', '肺活量(ml)'),
    ('sit_and_reach', '坐位体前屈(cm)'),
    ('run_50m', '50米(秒)'),
    ('standing_long_jump', '立定跳远(cm)'),
    ('strength', '一分钟仰卧起坐（女）/引体向上(男)（次）'),
    ('endurance', '800米跑（女）/1000米跑（男）'),
]

# 单项指标 key -> 体测记录里的标准分字段
SCORE_KEY_MAP = {
    'bmi': 'bmi_score',
    'vital_capacity': 'vital_score',
    'sit_and_reach': 'sitreach_score',
    'run_50m': 'run50_score',
    'standing_long_jump': 'jump_score',
    'strength': 'strength_score',
    'endurance': 'endurance_score',
}


def _default_year():
    return str(datetime.date.today().year)


def _year_options():
    """年份下拉选项：当前年份 ~ 2000（降序，默认选中当前年份）。"""
    current = datetime.date.today().year
    return [str(y) for y in range(current, 1999, -1)]


def _all_classes():
    return SchoolClass.query.order_by(SchoolClass.id).all()


def _f(v):
    if v is None:
        return None
    s = str(v).strip()
    if s == '':
        return None
    try:
        return float(s)
    except (TypeError, ValueError):
        return None


def _parse_endurance(v):
    """把「分.秒」耐力成绩转为秒。4.30 = 4分30秒 = 270秒；4.3 = 4分30秒；4 = 4分0秒。"""
    if v is None:
        return None
    s = str(v).strip()
    if s == '':
        return None
    try:
        if '.' in s:
            m, sec = s.split('.', 1)
            minutes = int(float(m))
            seconds = int(sec.ljust(2, '0')[:2])
        else:
            minutes = int(float(s))
            seconds = 0
        return minutes * 60 + seconds
    except (TypeError, ValueError):
        return None


def _strength_parts(gender, strength, endurance, m):
    if gender == '男':
        m['pull_up'] = strength if strength is not None else None
        m['run_1000m'] = endurance
    else:
        m['sit_up'] = strength if strength is not None else None
        m['run_800m'] = endurance
    return m


def _apply_scores(rec, gender, m):
    r = calculate(gender, m)
    rec.height_cm = m.get('height_cm')
    rec.weight_kg = m.get('weight_kg')
    rec.bmi = r['bmi']
    rec.vital_capacity = m.get('vital_capacity')
    rec.run_50m = m.get('run_50m')
    rec.sit_and_reach = m.get('sit_and_reach')
    rec.standing_long_jump = m.get('standing_long_jump')
    rec.pull_up = m.get('pull_up')
    rec.sit_up = m.get('sit_up')
    rec.run_1000m = m.get('run_1000m')
    rec.run_800m = m.get('run_800m')
    rec.bmi_score = r['bmi_score']
    rec.vital_score = r['vital_score']
    rec.run50_score = r['run50_score']
    rec.sitreach_score = r['sitreach_score']
    rec.jump_score = r['jump_score']
    rec.strength_score = r['strength_score']
    rec.endurance_score = r['endurance_score']
    rec.extra_score = 0.0
    rec.total_score = r['total_score']
    rec.grade = r['grade']
    return rec


def _records_map(cid, year):
    """返回 {student_id: FitnessRecord}（按班级+年份过滤）。"""
    q = FitnessRecord.query.filter(FitnessRecord.student_id.in_(
        db.session.query(Student.id).filter(Student.class_id == cid)))
    if year:
        q = q.filter(FitnessRecord.year == year)
    return {r.student_id: r for r in q.all()}


def _score_of(rec, item):
    key = SCORE_KEY_MAP.get(item)
    return getattr(rec, key, None) if key else None


@modules_bp.route('/fitness')
@login_required
def fitness():
    return redirect(url_for('modules.fitness_report'))


# ============ 1. 体测评分录入 ============

@modules_bp.route('/fitness/entry')
@login_required
def fitness_entry():
    classes = _all_classes()
    cid = request.args.get('class_id', type=int) or (classes[0].id if classes else None)
    year = request.args.get('year', '').strip() or _default_year()
    students = Student.query.filter_by(class_id=cid).order_by(Student.student_no).all() if cid else []
    return render_template('fitness/entry.html', active='fitness_entry',
                           classes=classes, cid=cid, year=year, students=students,
                           year_options=_year_options())


@modules_bp.route('/fitness/add', methods=['POST'])
@login_required
def fitness_add():
    student_id = request.form.get('student_id', type=int)
    year = request.form.get('year', '').strip() or _default_year()
    s = db.session.get(Student, student_id) if student_id else None
    if not s:
        flash('请选择学生', 'error')
        return redirect(url_for('modules.fitness_entry'))
    m = {'bmi': _f(request.form.get('bmi')),
         'vital_capacity': _f(request.form.get('vital_capacity')), 'run_50m': _f(request.form.get('run_50m')),
         'sit_and_reach': _f(request.form.get('sit_and_reach')), 'standing_long_jump': _f(request.form.get('standing_long_jump'))}
    m = _strength_parts(s.gender or '男', _f(request.form.get('strength')), _parse_endurance(request.form.get('endurance')), m)
    rec = FitnessRecord.query.filter_by(student_id=student_id, year=year).first()
    if not rec:
        rec = FitnessRecord(student_id=student_id, year=year)
        db.session.add(rec)
    _apply_scores(rec, s.gender or '男', m)
    db.session.commit()
    flash('体测数据已保存并自动评分', 'success')
    return redirect(url_for('modules.fitness_entry', class_id=s.class_id, year=year))


@modules_bp.route('/fitness/import', methods=['POST'])
@login_required
def fitness_import():
    cid = request.form.get('class_id', type=int)
    year = request.form.get('year', '').strip() or _default_year()
    if not cid:
        flash('请选择班级', 'error')
        return redirect(url_for('modules.fitness_entry'))
    f = request.files.get('file')
    if not f or not f.filename:
        flash('请选择要导入的文件', 'error')
        return redirect(url_for('modules.fitness_entry', class_id=cid, year=year))
    rows, errors = _parse_fitness_file(f)
    if errors:
        flash('导入校验失败：' + '；'.join(errors[:5]), 'error')
        return redirect(url_for('modules.fitness_entry', class_id=cid, year=year))
    students = {s.student_no: s for s in Student.query.filter_by(class_id=cid).all()}
    imported = 0
    for r in rows:
        s = students.get(r['no'])
        if not s:
            continue
        gender = r.get('gender') if r.get('gender') in ('男', '女') else (s.gender or '男')
        m = {'bmi': r.get('bmi'), 'vital_capacity': r.get('vital_capacity'),
             'run_50m': r.get('run_50m'), 'sit_and_reach': r.get('sit_and_reach'),
             'standing_long_jump': r.get('standing_long_jump')}
        m = _strength_parts(gender, r.get('strength'), r.get('endurance'), m)
        rec = FitnessRecord.query.filter_by(student_id=s.id, year=year).first()
        if not rec:
            rec = FitnessRecord(student_id=s.id, year=year)
            db.session.add(rec)
        _apply_scores(rec, gender, m)
        imported += 1
    db.session.commit()
    flash('成功导入 %d 条体测记录并自动评分' % imported, 'success')
    return redirect(url_for('modules.fitness_entry', class_id=cid, year=year))


@modules_bp.route('/fitness/<int:rid>/delete', methods=['POST'])
@login_required
def fitness_delete(rid):
    r = db.get_or_404(FitnessRecord, rid)
    class_id = db.session.get(Student, r.student_id).class_id if r.student_id else None
    year = r.year
    db.session.delete(r)
    db.session.commit()
    flash('已删除体测记录', 'success')
    return redirect(request.referrer or url_for('modules.fitness_query', class_id=class_id, year=year))


@modules_bp.route('/fitness/template')
@login_required
def fitness_template():
    buf = _build_fitness_template()
    return send_file(buf, as_attachment=True, download_name='体测成绩导入模板.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ============ 2. 体质数据查询 ============

@modules_bp.route('/fitness/query')
@login_required
def fitness_query():
    classes = _all_classes()
    cid = request.args.get('class_id', type=int)
    year = request.args.get('year', '').strip()
    keyword = request.args.get('keyword', '').strip()
    rows = _query_rows(cid, year, keyword)
    class_map = {c.id: c.name for c in classes}
    return render_template('fitness/query.html', active='fitness_query',
                           classes=classes, cid=cid, year=year, keyword=keyword,
                           rows=rows, class_map=class_map, year_options=_year_options())


def _query_rows(cid, year, keyword):
    # 内连接：仅返回有体测记录的学生；年份作为 WHERE 过滤
    q = db.session.query(Student, FitnessRecord).join(
        FitnessRecord, FitnessRecord.student_id == Student.id)
    if cid:
        q = q.filter(Student.class_id == cid)
    if year:
        q = q.filter(FitnessRecord.year == year)
    if keyword:
        q = q.filter(db.or_(Student.name.like('%%%s%%' % keyword), Student.student_no.like('%%%s%%' % keyword)))
    q = q.order_by(Student.class_id, Student.student_no)
    return q.limit(500).all()


# ============ 3. 体质统计报表 ============

@modules_bp.route('/fitness/report')
@login_required
def fitness_report():
    classes = _all_classes()
    grades = sorted({c.grade for c in classes if c.grade})
    cid = request.args.get('class_id', type=int) or (classes[0].id if classes else None)
    grade = request.args.get('grade', '').strip()
    year = request.args.get('year', '').strip() or _default_year()
    item = request.args.get('item', '').strip()

    # 班级过滤（年级 + 班级）
    cls_list = classes
    if grade:
        cls_list = [c for c in classes if c.grade == grade]
    if cid and (not grade or cid in [c.id for c in cls_list]):
        cls_list = [c for c in cls_list if c.id == cid]

    students = []
    for c in cls_list:
        students += Student.query.filter_by(class_id=c.id).order_by(Student.student_no).all()
    records = _records_map(cid, year) if cid else {}

    stats = _compute_stats(students, records, item)
    year_trend = _year_trend(cid, item) if cid else []
    item_dist = _item_dist(students, records, item) if item else []
    return render_template('fitness/report.html', active='fitness_report',
                           classes=classes, grades=grades, cid=cid, grade=grade, year=year,
                           item=item, item_labels=ITEM_LABELS,
                           students=students, records=records, stats=stats,
                           year_trend=year_trend, item_dist=item_dist,
                           year_options=_year_options())


def _grade_of(score):
    if score >= 90:
        return '优秀'
    if score >= 80:
        return '良好'
    if score >= 60:
        return '合格'
    return '不及格'


def _compute_stats(students, records, item=''):
    """统计（可选按单项指标）：item 为空时按总分统计，否则按该指标标准分统计。"""
    score_key = SCORE_KEY_MAP.get(item) if item else 'total_score'
    dist = {'优秀': 0, '良好': 0, '合格': 0, '不及格': 0}
    total = 0.0
    n = 0
    for s in students:
        r = records.get(s.id)
        if not r:
            continue
        score = getattr(r, score_key, None)
        if score is None:
            continue
        n += 1
        total += score
        dist[_grade_of(score)] += 1
    avg = round(total / n, 1) if n else 0
    return {'dist': dist, 'n': n, 'total_students': len(students), 'avg': avg}


def _year_trend(cid, item=''):
    if not cid:
        return []
    score_col = getattr(FitnessRecord, SCORE_KEY_MAP[item]) if item else FitnessRecord.total_score
    q = db.session.query(FitnessRecord.year, db.func.avg(score_col)).filter(
        FitnessRecord.student_id.in_(db.session.query(Student.id).filter(Student.class_id == cid)))
    q = q.group_by(FitnessRecord.year).order_by(FitnessRecord.year)
    return [(y, round(avg, 1) if avg else 0) for y, avg in q.all()]


def _item_dist(students, records, item):
    key = SCORE_KEY_MAP.get(item)
    if not key:
        return []
    buckets = {'<60': 0, '60-69': 0, '70-79': 0, '80-89': 0, '≥90': 0}
    for s in students:
        r = records.get(s.id)
        if not r:
            continue
        v = getattr(r, key, None)
        if v is None:
            continue
        if v < 60:
            buckets['<60'] += 1
        elif v < 70:
            buckets['60-69'] += 1
        elif v < 80:
            buckets['70-79'] += 1
        elif v < 90:
            buckets['80-89'] += 1
        else:
            buckets['≥90'] += 1
    return buckets


# ============ 4. 体质预警名单 ============

@modules_bp.route('/fitness/warnings')
@login_required
def fitness_warnings():
    classes = _all_classes()
    cid = request.args.get('class_id', type=int)
    year = request.args.get('year', '').strip() or _default_year()
    warn = _build_warnings(cid, year)
    return render_template('fitness/warnings.html', active='fitness_warnings',
                           classes=classes, cid=cid, year=year, warn=warn,
                           year_options=_year_options())


def _build_warnings(cid, year):
    cls_list = [c for c in _all_classes() if not cid or c.id == cid]
    result = []
    for c in cls_list:
        students = Student.query.filter_by(class_id=c.id).order_by(Student.student_no).all()
        records = _records_map(c.id, year)
        for s in students:
            r = records.get(s.id)
            if not r or r.total_score is None:
                continue
            weak = []
            for key, label in ITEM_LABELS:
                sc = _score_of(r, key)
                if sc is not None and sc < 60:
                    weak.append(label)
            level = 'high' if (r.grade == '不及格' or len(weak) >= 3) else ('mid' if (r.grade == '合格' or weak) else None)
            if level:
                reason = []
                if r.grade == '不及格':
                    reason.append('总分不及格')
                if weak:
                    reason.append('单项薄弱：' + '、'.join(weak))
                result.append({'student': s, 'record': r, 'class': c, 'weak': weak,
                               'level': level, 'reason': '；'.join(reason)})
    result.sort(key=lambda x: (x['record'].total_score or 0))
    return result


# ============ 5. 体质报告导出 ============

@modules_bp.route('/fitness/export')
@login_required
def fitness_export():
    classes = _all_classes()
    cid = request.args.get('class_id', type=int) or (classes[0].id if classes else None)
    year = request.args.get('year', '').strip() or _default_year()
    return render_template('fitness/export.html', active='fitness_export',
                           classes=classes, cid=cid, year=year, year_options=_year_options())


@modules_bp.route('/fitness/export/excel')
@login_required
def fitness_export_excel():
    cid = request.args.get('class_id', type=int)
    year = request.args.get('year', '').strip() or _default_year()
    if not cid:
        return '请选择班级', 400
    cls = db.session.get(SchoolClass, cid)
    students = Student.query.filter_by(class_id=cid).order_by(Student.student_no).all()
    records = _records_map(cid, year)
    buf = _build_excel(cls, students, records, year)
    return send_file(buf, as_attachment=True,
                     download_name='%s-体质报告-%s.xlsx' % (cls.name if cls else '班级', year),
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@modules_bp.route('/fitness/export/pdf')
@login_required
def fitness_report_pdf():
    """打印友好的 HTML 报告（浏览器另存为 PDF）。"""
    cid = request.args.get('class_id', type=int)
    year = request.args.get('year', '').strip() or _default_year()
    item = request.args.get('item', '').strip()
    item_label = dict(ITEM_LABELS).get(item, '') if item else ''
    cls = db.session.get(SchoolClass, cid) if cid else None
    students = Student.query.filter_by(class_id=cid).order_by(Student.student_no).all() if cid else []
    records = _records_map(cid, year)
    stats = _compute_stats(students, records, item)
    return render_template('fitness/report_pdf.html', cls=cls, year=year, item_label=item_label,
                           students=students, records=records, stats=stats)


# ============ 工具函数 ============

def _parse_fitness_file(file):
    """解析体测 Excel/CSV，返回 (rows, errors)。errors 为「第X行 原因」列表。"""
    filename = (file.filename or '').lower()
    try:
        if filename.endswith(('.xlsx', '.xls')):
            from openpyxl import load_workbook
            wb = load_workbook(file.stream, read_only=True, data_only=True)
            raw = [[('' if c is None else str(c).strip()) for c in row] for row in wb.active.iter_rows(values_only=True)]
        elif filename.endswith('.csv'):
            text = None
            for enc in ('utf-8-sig', 'gbk', 'utf-8'):
                try:
                    text = file.stream.read().decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            if text is None:
                return [], ['CSV 编码无法识别，请另存为 UTF-8 或 GBK']
            raw = [r for r in csv.reader(io.StringIO(text))]
        else:
            return [], ['仅支持 .xlsx / .xls / .csv 文件']
    except Exception as e:
        return [], ['文件解析失败：%s' % e]
    if not raw:
        return [], ['文件为空']
    header = [str(h).replace(' ', '').replace('（', '(').replace('）', ')') for h in raw[0]]
    col = {}
    for i, h in enumerate(header):
        hl = h.lower()
        if 'no' not in col and ('学号' in h or 'student_no' in hl):
            col['no'] = i
        elif 'name' not in col and '姓名' in h:
            col['name'] = i
        elif 'gender' not in col and '性别' in h:
            col['gender'] = i
        elif 'grade' not in col and '年级' in h:
            col['grade'] = i
        elif 'college' not in col and '学院' in h:
            col['college'] = i
        elif 'class' not in col and '班级' in h:
            col['class'] = i
        elif 'bmi' not in col and ('体重指数' in h or 'bmi' in hl):
            col['bmi'] = i
        elif 'vital' not in col and '肺活量' in h:
            col['vital'] = i
        elif 'sitreach' not in col and '坐位体前屈' in h:
            col['sitreach'] = i
        elif 'run50' not in col and '50米' in h:
            col['run50'] = i
        elif 'jump' not in col and '立定跳远' in h:
            col['jump'] = i
        elif 'strength' not in col and ('仰卧起坐' in h or '引体向上' in h or '力量' in h):
            col['strength'] = i
        elif 'endurance' not in col and ('800米' in h or '1000米' in h or '耐力' in h):
            col['endurance'] = i
    if 'no' not in col:
        return [], ['未找到「学号」列，请检查表头']
    rows, errors = [], []
    for idx, r in enumerate(raw[1:], start=2):
        if not r or all(not c for c in r):
            continue
        def cell(key):
            i = col.get(key)
            return str(r[i]).strip() if i is not None and i < len(r) else ''
        def fcell(key):
            i = col.get(key)
            return _f(r[i]) if i is not None and i < len(r) else None
        no = cell('no')
        name = cell('name')
        gender = cell('gender')
        grade = cell('grade')
        cls = cell('class')
        bmi = fcell('bmi')
        vital = fcell('vital')
        sitreach = fcell('sitreach')
        run50 = fcell('run50')
        jump = fcell('jump')
        strength = fcell('strength')
        endurance = _parse_endurance(cell('endurance'))
        row_errors = []
        if not no:
            row_errors.append('学号为空')
        if not name:
            row_errors.append('姓名为空')
        if gender not in ('男', '女'):
            row_errors.append('性别填写错误，仅允许男/女')
        if not grade:
            row_errors.append('年级为空')
        if not cls:
            row_errors.append('班级为空')
        if row_errors:
            for e in row_errors:
                errors.append('第%d行，%s' % (idx, e))
            continue
        rows.append({'no': no, 'gender': gender, 'bmi': bmi, 'vital_capacity': vital,
                     'sit_and_reach': sitreach, 'run_50m': run50, 'standing_long_jump': jump,
                     'strength': strength, 'endurance': endurance})
    return rows, errors


def _build_fitness_template():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = '体测成绩'
    headers = ['学号', '姓名', '性别', '年级', '学院', '班级', '体重指数（BMI）', '肺活量(ml)',
               '坐位体前屈(cm)', '50米(秒)', '立定跳远(cm)', '一分钟仰卧起坐（女）/引体向上(男)（次）', '800米跑（女）/1000米跑（男）']
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h).font = Font(bold=True)
    notes = ['必填', '必填', '仅填男/女', '必填', '选填', '必填', '单位：无',
             '单位：ml', '单位：cm', '单位：秒', '单位：cm',
             '男：引体向上次数；女：一分钟仰卧起坐次数', '男：1000米「分.秒」如4.30；女：800米「分.秒」如4.30']
    for i, n in enumerate(notes, start=1):
        ws.cell(row=2, column=i, value=n).font = Font(color='888888', size=10)
    for i, w in enumerate((12, 10, 8, 10, 14, 12, 14, 12, 14, 10, 12, 24, 22), start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws2 = wb.create_sheet('使用说明')
    for i, line in enumerate(['体测成绩导入说明',
                              '1. 学号/姓名/性别/年级/班级为必填，性别仅允许填写「男」或「女」。',
                              '2. 力量项目：男生填引体向上次数，女生填一分钟仰卧起坐次数。',
                              '3. 耐力项目：男生填1000米成绩，女生填800米成绩；按「分.秒」填写，如 4.30 表示4分30秒。',
                              '4. 各项目按单位填写即可，不做数值范围限制。'], start=1):
        ws2.cell(row=i, column=1, value=line)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _build_excel(cls, students, records, year):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = '体质报告'
    ws.cell(row=1, column=1, value='%s 体质健康报告（%s）' % (
        cls.name if cls else '', year)).font = Font(bold=True, size=14)
    headers = ['学号', '姓名', '性别', '总分', '等级', '肺活量', '50米', '立定跳远', '力量', '耐力跑']
    for i, h in enumerate(headers, start=1):
        ws.cell(row=3, column=i, value=h).font = Font(bold=True)
    row = 4
    for s in students:
        r = records.get(s.id)
        if not r:
            continue
        vals = [s.student_no, s.name, s.gender, r.total_score, r.grade,
                r.vital_capacity, r.run_50m, r.standing_long_jump,
                (r.pull_up if s.gender == '男' else r.sit_up), (r.run_1000m if s.gender == '男' else r.run_800m)]
        for i, v in enumerate(vals, start=1):
            ws.cell(row=row, column=i, value=v if v is not None else '')
        row += 1
    for i, w in enumerate((12, 10, 8, 8, 8, 10, 8, 10, 8, 10), start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@modules_bp.route('/action')
@login_required
def action():
    """AI 动作识别入口：重定向到「上传视频」。"""
    return redirect(url_for('modules.action_upload'))


# ============ 共享工具 ============

def _student_name(a):
    """解析诊断记录的学生姓名（未关联学生时返回空字符串）。"""
    if a.student_id:
        s = db.session.get(Student, a.student_id)
        if s:
            return s.name or ''
    return ''


def _first_landmarks(a):
    """取诊断记录首帧关键点（供弹窗绘制骨骼），无则返回 None。"""
    try:
        pose = json.loads(a.pose_data) if a.pose_data else None
    except Exception:
        pose = None
    if pose:
        for f in pose:
            if isinstance(f, dict) and f.get('landmarks'):
                return f['landmarks']
    return None


def _split_names(students):
    """把 AI 返回的学生名单拆成单个姓名（兼容顿号/逗号/分号分隔，及整体字符串）。"""
    if isinstance(students, str):
        students = [students]
    names = []
    for s in students or []:
        if not s:
            continue
        for part in re.split(r'[、，,;；/]', str(s)):
            part = part.strip()
            if part:
                names.append(part)
    return names


def _action_detail(a):
    """组装弹窗/报告共用的单条诊断详情数据。"""
    return {
        'id': a.id,
        'student': _student_name(a) or '未指定学生',
        'sport_type': a.sport_type or '',
        'technique_action': a.technique_action or '',
        'video_path': a.video_path or '',
        'result': a.result or '',
        'metrics': json.loads(a.metrics_data) if a.metrics_data else None,
        'landmarks': _first_landmarks(a),
        'status': a.status or 'success',
        'error': a.error or '',
        'created_at': a.created_at.strftime('%Y-%m-%d %H:%M') if a.created_at else '',
    }


def _cleanup_frames(frame_paths):
    """删除抽帧图片（失败清理用）。"""
    for p in frame_paths or []:
        try:
            os.remove(p)
        except Exception:
            pass


def _run_pose(tmp_video_path):
    """姿态分析：返回 (pose_status, pose_msg, pose_data, metrics_data, key_frames)。"""
    if not MP_AVAILABLE:
        return ('not_install',
                'MediaPipe 未安装，暂无关节参数。请关闭 Flask 服务后执行 pip install mediapipe 完成安装。',
                None, None, None)
    try:
        pose_data, metrics_data, key_frames = analyze_pose(tmp_video_path)
        return 'success', '', pose_data, metrics_data, key_frames
    except Exception as e:
        msg = str(e) or ''
        if '未检测到人体' in msg or '关键点' in msg:
            return 'no_human', '未检测到视频中的人体，无法生成关节参数。', None, None, None
        return 'no_human', '姿态分析失败：%s' % msg, None, None, None


def _analyze_record(rec, tmp_video_path, ext):
    """执行单条诊断分析并更新记录（单条与批量共用）。失败抛异常并清理帧。"""
    frame_paths = extract_frames(tmp_video_path)
    if not frame_paths:
        raise ValueError('视频抽帧失败，请确认文件是有效视频')
    try:
        result = analyze_action(frame_paths, rec.sport_type or '', _student_name(rec),
                                rec.technique_action or '')
        result = (result or '').strip()
        if not result:
            raise ValueError('视觉模型未返回诊断内容')
        pose_status, pose_msg, pose_data, metrics_data, key_frames = _run_pose(tmp_video_path)
        # 归档视频到 uploads
        filename = 'action_%s%s' % (uuid.uuid4().hex[:10], ext)
        final_path = os.path.join(UPLOAD_DIR, filename)
        os.replace(tmp_video_path, final_path)
        rec.video_path = '/static/uploads/' + filename
        rec.frames = json.dumps(frame_paths)
        rec.result = result
        rec.pose_data = json.dumps(pose_data) if pose_data else None
        rec.metrics_data = json.dumps(metrics_data) if metrics_data else None
        rec.key_frames = json.dumps(key_frames) if key_frames else None
        rec.status = pose_status if pose_status != 'success' else 'success'
        rec.error = pose_msg if pose_status != 'success' else ''
        return rec
    except Exception:
        _cleanup_frames(frame_paths)
        raise


def _parse_student_from_filename(filename):
    """从文件名解析学生姓名：去扩展名、去序号前缀与括号说明。"""
    name = os.path.splitext(filename)[0].strip()
    name = re.sub(r'^[\d\s\-_]+', '', name)
    name = re.sub(r'[（(].*?[)）]', '', name).strip()
    return name


def _match_student(name, class_id=None):
    """按姓名匹配学生库，返回 Student 或 None。"""
    if not name:
        return None
    q = Student.query.filter_by(name=name)
    if class_id:
        q = q.filter_by(class_id=class_id)
    s = q.first()
    if s:
        return s
    return Student.query.filter_by(name=name).first()


def _class_name(cid):
    if not cid:
        return '全部班级'
    c = db.session.get(SchoolClass, cid)
    return c.name if c else '全部班级'


# ============ 1. 上传视频 ============

@modules_bp.route('/action/upload')
@login_required
def action_upload():
    classes = _all_classes()
    students = Student.query.order_by(Student.class_id, Student.student_no).all()
    class_map = {c.id: c.name for c in classes}
    return render_template('action/upload.html', active='action_upload',
                           classes=classes, students=students, class_map=class_map,
                           sport_types=SPORT_ITEMS, techniques=SPORT_TECHNIQUES,
                           mediapipe_installed=MP_AVAILABLE)


@modules_bp.route('/action/analyze', methods=['POST'])
@login_required
def action_analyze():
    """单条视频诊断（同步）。"""
    cid = request.form.get('class_id', type=int)
    student_id = request.form.get('student_id', type=int)
    sport_type = request.form.get('sport_type', '').strip()
    technique_action = request.form.get('technique_action', '').strip()
    f = request.files.get('video')
    if not f or not f.filename:
        flash('请选择要上传的视频', 'error')
        return redirect(url_for('modules.action_upload'))
    ext = os.path.splitext(f.filename)[1].lower()
    if ext not in ('.mp4', '.mov', '.avi', '.mkv', '.webm'):
        flash('仅支持 mp4 / mov / avi / mkv / webm 视频', 'error')
        return redirect(url_for('modules.action_upload'))
    if not sport_type:
        flash('请选择运动项目', 'error')
        return redirect(url_for('modules.action_upload'))
    if technique_action == '未指定':
        technique_action = ''
    # 班级：未显式指定时从学生所属班级推导
    student_name = ''
    if student_id:
        s = db.session.get(Student, student_id)
        student_name = s.name if s else ''
        if not cid and s and s.class_id:
            cid = s.class_id
    # 存到 ASCII 临时路径（OpenCV 抽帧不支持中文路径）
    import tempfile
    tmp_path = os.path.join(tempfile.gettempdir(), 'action_%s%s' % (uuid.uuid4().hex[:10], ext))
    f.save(tmp_path)
    rec = ActionAnalysis(student_id=student_id, class_id=cid, sport_type=sport_type,
                         technique_action=technique_action, status='processing')
    db.session.add(rec)
    db.session.commit()
    try:
        _analyze_record(rec, tmp_path, ext)
        db.session.commit()
        flash('动作识别完成，已生成纠错建议', 'success')
    except Exception as e:
        db.session.rollback()
        db.session.delete(rec)
        db.session.commit()
        try:
            os.remove(tmp_path)
        except Exception:
            pass
        flash('动作识别失败：%s' % e, 'error')
    return redirect(url_for('modules.action_history', class_id=cid, sport_type=sport_type))


# 批量任务内存队列
_ACTION_BATCHES = {}
_ACTION_BATCH_LOCK = threading.Lock()


def _run_batch(batch_id, app):
    with app.app_context():
        job = _ACTION_BATCHES.get(batch_id)
        if not job:
            return
        for item in job['items']:
            rec = db.session.get(ActionAnalysis, item['analysis_id'])
            if not rec:
                job['done'] += 1
                continue
            rec.status = 'processing'
            db.session.commit()
            try:
                _analyze_record(rec, item['tmp_path'], item['ext'])
                db.session.commit()
                job['success'] += 1
            except Exception as e:
                try:
                    db.session.rollback()
                except Exception:
                    pass
                try:
                    os.remove(item['tmp_path'])
                except Exception:
                    pass
                rec = db.session.get(ActionAnalysis, item['analysis_id'])
                if rec:
                    rec.status = 'failed'
                    rec.error = str(e)
                    db.session.commit()
                job['failed'].append({'student': item['student'], 'filename': item['filename'],
                                      'error': str(e)})
            job['done'] += 1
        job['running'] = False


@modules_bp.route('/action/batch-upload', methods=['POST'])
@login_required
def action_batch_upload():
    """批量上传：异步队列处理，返回 batch_id 供前端轮询进度。"""
    cid = request.form.get('class_id', type=int)
    sport_type = request.form.get('sport_type', '').strip()
    technique_action = request.form.get('technique_action', '').strip()
    files = request.files.getlist('videos')
    if not sport_type:
        return jsonify({'error': '请选择运动项目'}), 400
    if technique_action == '未指定':
        technique_action = ''
    if not files or not files[0].filename:
        return jsonify({'error': '请选择要上传的视频'}), 400

    import tempfile
    batch_id = uuid.uuid4().hex[:12]
    items = []
    total = 0
    for f in files:
        if not f.filename:
            continue
        ext = os.path.splitext(f.filename)[1].lower()
        if ext not in ('.mp4', '.mov', '.avi', '.mkv', '.webm'):
            continue
        name = _parse_student_from_filename(f.filename)
        matched = _match_student(name, cid) if name else None
        tmp_path = os.path.join(tempfile.gettempdir(), 'action_%s_%d%s' % (batch_id, total, ext))
        f.save(tmp_path)
        rec = ActionAnalysis(student_id=matched.id if matched else None,
                             class_id=(matched.class_id if matched else cid),
                             sport_type=sport_type, technique_action=technique_action,
                             status='pending',
                             error=('未匹配学生：文件名「%s」未在班级中找到对应学生' % name) if (name and not matched) else '')
        db.session.add(rec)
        db.session.commit()
        items.append({'analysis_id': rec.id, 'tmp_path': tmp_path, 'ext': ext,
                      'student': (matched.name if matched else (name or '未匹配学生')),
                      'filename': f.filename, 'matched': bool(matched)})
        total += 1

    if not items:
        return jsonify({'error': '没有可处理的视频文件（仅支持 mp4/mov/avi/mkv/webm）'}), 400

    with _ACTION_BATCH_LOCK:
        _ACTION_BATCHES[batch_id] = {
            'batch_id': batch_id, 'total': total, 'done': 0, 'success': 0,
            'failed': [], 'items': items, 'running': True,
        }
    app_obj = current_app._get_current_object()
    t = threading.Thread(target=_run_batch, args=(batch_id, app_obj))
    t.daemon = True
    t.start()
    return jsonify({'batch_id': batch_id, 'total': total})


@modules_bp.route('/action/batch/status')
@login_required
def action_batch_status():
    """批量任务进度查询。"""
    batch_id = request.args.get('batch_id', '')
    job = _ACTION_BATCHES.get(batch_id)
    if not job:
        return jsonify({'error': '未找到批量任务'})
    return jsonify({
        'batch_id': batch_id, 'total': job['total'], 'done': job['done'],
        'success': job['success'], 'running': job['running'],
        'failed': job['failed'], 'failed_count': len(job['failed']),
    })


# ============ 2. 识别历史 ============

@modules_bp.route('/action/history')
@login_required
def action_history():
    classes = _all_classes()
    cid = request.args.get('class_id', type=int)
    sport_type = request.args.get('sport_type', '').strip()
    student_id = request.args.get('student_id', type=int)
    date_from = request.args.get('date_from', '').strip()
    date_to = request.args.get('date_to', '').strip()
    q = ActionAnalysis.query
    if cid:
        q = q.filter(ActionAnalysis.class_id == cid)
    if sport_type:
        q = q.filter(ActionAnalysis.sport_type == sport_type)
    if student_id:
        q = q.filter(ActionAnalysis.student_id == student_id)
    if date_from:
        q = q.filter(ActionAnalysis.created_at >= date_from + ' 00:00:00')
    if date_to:
        q = q.filter(ActionAnalysis.created_at <= date_to + ' 23:59:59')
    analyses = q.order_by(ActionAnalysis.id.desc()).limit(500).all()
    students = Student.query.order_by(Student.class_id, Student.student_no).all()
    student_map = {s.id: s.name for s in Student.query.all()}
    class_map = {c.id: c.name for c in classes}
    return render_template('action/history.html', active='action_history',
                           classes=classes, cid=cid, sport_type=sport_type,
                           student_id=student_id, date_from=date_from, date_to=date_to,
                           sport_types=SPORT_ITEMS, analyses=analyses,
                           student_map=student_map, students=students, class_map=class_map,
                           status_labels=ACTION_STATUS_LABELS)


# ============ 3. 常见错误汇总 ============

@modules_bp.route('/action/summary')
@login_required
def action_summary():
    """常见错误汇总页面。"""
    classes = _all_classes()
    return render_template('action/summary.html', active='action_summary',
                           classes=classes, sport_types=SPORT_ITEMS, techniques=SPORT_TECHNIQUES)


def _query_summary(cid, sport_type, technique_action):
    """班级常见错误汇总查询 + AI 归类（专项训练方案与错误汇总共用的同一套数据源）。

    返回 (data, error)：成功时 data 为汇总结果 dict、error 为 None；失败时 data 为 None、error 为提示文案。
    """
    q = ActionAnalysis.query
    if cid:
        q = q.filter(ActionAnalysis.class_id == cid)
    if sport_type:
        q = q.filter(ActionAnalysis.sport_type == sport_type)
    if technique_action:
        q = q.filter(ActionAnalysis.technique_action == technique_action)
    records = q.order_by(ActionAnalysis.id.desc()).all()
    if not records:
        has_any = ActionAnalysis.query.count() > 0
        return None, ('当前筛选条件下没有找到对应的诊断记录，请更换筛选条件' if has_any
                      else '暂无诊断记录，请先上传训练视频并完成动作识别')
    # 每个学生取最近一次诊断（records 已按 id 倒序，首次即最近）
    name_to_analysis = {}
    for a in records:
        name = _student_name(a)
        if not name:
            continue
        name_to_analysis.setdefault(name, a)
    if not name_to_analysis:
        return None, '诊断记录均未关联学生，无法按学生汇总'

    items = [{'student': name, 'text': a.result or ''} for name, a in name_to_analysis.items()]
    try:
        groups = summarize_class_errors(items)
    except Exception as e:
        return None, 'AI 归类失败：%s' % e

    result_groups = []
    for g in groups:
        seen = set()
        students = []
        techs = []
        for nm in _split_names(g.get('students')):
            if nm in seen:
                continue
            seen.add(nm)
            a = name_to_analysis.get(nm)
            students.append({'name': nm, 'analysis_id': a.id if a else None})
            if a and a.technique_action:
                techs.append(a.technique_action)
        if not students:
            continue
        unique_techs = []
        for t in techs:
            if t not in unique_techs:
                unique_techs.append(t)
        if not unique_techs:
            unique_techs = ['未识别技术动作']
        result_groups.append({
            'error_name': g.get('error_name', ''),
            'count': len(students),
            'students': students,
            'techniques': unique_techs,
            'description': g.get('description', ''),
            'suggestion': g.get('suggestion', ''),
        })
    if not result_groups:
        return None, 'AI 未返回有效的错误归类结果，请重试'

    analyses = {}
    for a in name_to_analysis.values():
        analyses[a.id] = {
            'video_path': a.video_path or '',
            'result': a.result or '',
            'metrics': json.loads(a.metrics_data) if a.metrics_data else None,
            'landmarks': _first_landmarks(a),
            'technique_action': a.technique_action or '未识别技术动作',
        }
    data = {
        'class_name': _class_name(cid),
        'sport_name': sport_type or '全部项目',
        'technique_name': technique_action or '全部技术动作',
        'generated_at': datetime.datetime.now().strftime('%Y-%m-%d %H:%M'),
        'total_students': len(name_to_analysis),
        'groups': result_groups,
        'analyses': analyses,
    }
    return data, None


@modules_bp.route('/action/summary/data')
@login_required
def action_summary_data():
    """班级常见错误汇总数据接口（count_only=1 时仅做轻量预检，不触发 AI 归类）。"""
    cid = request.args.get('class_id', type=int)
    sport_type = request.args.get('sport_type', '').strip()
    technique_action = request.args.get('technique_action', '').strip()
    if request.args.get('count_only') == '1':
        q = ActionAnalysis.query
        if cid:
            q = q.filter(ActionAnalysis.class_id == cid)
        if sport_type:
            q = q.filter(ActionAnalysis.sport_type == sport_type)
        if technique_action:
            q = q.filter(ActionAnalysis.technique_action == technique_action)
        has_data = any(_student_name(a) for a in q.all())
        return jsonify({'has_data': has_data})
    data, error = _query_summary(cid, sport_type, technique_action)
    if error:
        return jsonify({'error': error})
    return jsonify(data)


@modules_bp.route('/action/summary/pdf', methods=['POST'])
@login_required
def action_summary_pdf():
    """班级常见错误汇总 PDF（打印友好 HTML，按错误类型分组）。"""
    raw = request.form.get('data', '')
    try:
        data = json.loads(raw) if raw else {}
    except Exception:
        data = {}
    groups = []
    for g in data.get('groups') or []:
        if not isinstance(g, dict) or not g.get('error_name'):
            continue
        names = _split_names([s.get('name', '') if isinstance(s, dict) else s
                              for s in (g.get('students') or [])])
        groups.append({
            'error_name': str(g.get('error_name', '')).strip(),
            'count': int(g.get('count') or 0),
            'students': names,
            'techniques': [str(t) for t in (g.get('techniques') or [])],
            'description': str(g.get('description', '')),
            'suggestion': str(g.get('suggestion', '')),
        })
    total_students = int(data.get('total_students') or 0)
    involved = sum(g['count'] for g in groups)
    return render_template('action/summary_pdf.html',
                           title=str(data.get('title') or '班级常见错误汇总'),
                           subtitle=str(data.get('subtitle') or ''),
                           total_students=total_students, involved=involved, groups=groups)


# ============ 4. 动作诊断报告库 ============

@modules_bp.route('/action/reports')
@login_required
def action_reports():
    classes = _all_classes()
    cid = request.args.get('class_id', type=int)
    sport_type = request.args.get('sport_type', '').strip()
    student_id = request.args.get('student_id', type=int)
    year = request.args.get('year', '').strip()
    q = ActionAnalysis.query
    if cid:
        q = q.filter(ActionAnalysis.class_id == cid)
    if sport_type:
        q = q.filter(ActionAnalysis.sport_type == sport_type)
    if student_id:
        q = q.filter(ActionAnalysis.student_id == student_id)
    analyses = q.order_by(ActionAnalysis.id.desc()).all()
    if year:
        analyses = [a for a in analyses if a.created_at and str(a.created_at.year) == year]
    analyses = analyses[:500]
    students = Student.query.order_by(Student.class_id, Student.student_no).all()
    student_map = {s.id: s.name for s in students}
    class_map = {c.id: c.name for c in classes}
    years = sorted({str(a.created_at.year) for a in ActionAnalysis.query.all()
                    if a.created_at and a.created_at.year}, reverse=True)
    return render_template('action/reports.html', active='action_reports',
                           classes=classes, cid=cid, sport_type=sport_type, student_id=student_id,
                           year=year, sport_types=SPORT_ITEMS, analyses=analyses,
                           student_map=student_map, students=students, class_map=class_map,
                           years=years, status_labels=ACTION_STATUS_LABELS)


@modules_bp.route('/action/report/<int:aid>')
@login_required
def action_report_pdf(aid):
    """单人诊断报告（打印友好 HTML，浏览器另存为 PDF）。"""
    a = db.get_or_404(ActionAnalysis, aid)
    return render_template('action/report_pdf.html', d=_action_detail(a))


@modules_bp.route('/action/reports/pdf', methods=['POST'])
@login_required
def action_reports_pdf():
    """批量导出多个学生的单人报告（打印友好 HTML）。"""
    ids = request.form.getlist('ids', type=int)
    records = []
    for aid in ids:
        a = db.session.get(ActionAnalysis, aid)
        if a:
            records.append(_action_detail(a))
    return render_template('action/reports_pdf.html', records=records)


# ============ 5. 专项训练方案 ============

@modules_bp.route('/action/training')
@login_required
def action_training():
    classes = _all_classes()
    return render_template('action/training.html', active='action_training',
                           classes=classes, sport_types=SPORT_ITEMS, techniques=SPORT_TECHNIQUES)


@modules_bp.route('/action/training/generate', methods=['POST'])
@login_required
def action_training_generate():
    """基于常见错误汇总生成班级专项训练方案（复用与错误汇总完全一致的查询逻辑）。"""
    data = request.get_json(silent=True) or {}
    cid = data.get('class_id')
    cid = int(cid) if cid not in (None, '') else None
    sport_type = data.get('sport_type', '').strip()
    technique_action = data.get('technique_action', '').strip()
    # 复用与「常见错误汇总」完全一致的查询逻辑，保证筛选规则对齐
    summary, error = _query_summary(cid, sport_type, technique_action)
    if error:
        return jsonify({'error': error})
    groups = summary['groups']
    lines = []
    for i, g in enumerate(groups, 1):
        lines.append('%d. %s（%d人）：%s' % (
            i, g.get('error_name', ''), int(g.get('count') or 0), g.get('suggestion', '')))
    prompt = (
        '你是高校体育专项训练方案设计助手。以下是「%s」的班级常见动作错误清单：\n\n%s\n\n'
        '请针对这些共性错误，生成一份班级专项训练方案，严格只输出一个 JSON 对象，字段：'
        'title(方案标题，简短)、goals(训练目标，一段话)、'
        'practices(数组，每项含 error_name 与 practice 两个字段，practice 为该错误对应的针对性练习建议)、'
        'schedule(课时安排建议，一段话)、points(练习要点，一段话)。'
        '所有字段内容请使用简体中文输出，不要使用英文。不要输出 JSON 以外的任何文字。' % (
            sport_type or '全部项目', '\n'.join(lines))
    )
    try:
        client = get_client('deepseek')
        resp = client.chat.completions.create(
            model=AGGREGATION_MODEL,
            messages=[{'role': 'user', 'content': prompt}],
            temperature=0.4, max_tokens=2500,
        )
        text = (resp.choices[0].message.content or '').strip() if resp.choices else ''
    except Exception:
        return jsonify({'error': 'AI生成训练方案失败，请稍后重试'})
    text = text.strip()
    if text.startswith('```'):
        text = text.strip('`')
        if text.lower().startswith('json'):
            text = text[4:]
        text = text.strip()
    try:
        plan = json.loads(text)
    except Exception:
        return jsonify({'error': 'AI生成训练方案失败，请稍后重试'})
    if not isinstance(plan, dict):
        return jsonify({'error': 'AI生成训练方案失败，请稍后重试'})
    plan['generated_at'] = datetime.datetime.now().strftime('%Y-%m-%d %H:%M')
    return jsonify({'plan': plan})


@modules_bp.route('/action/training/save', methods=['POST'])
@login_required
def action_training_save():
    """保存训练方案。"""
    data = request.get_json(silent=True) or {}
    cid = data.get('class_id', type=int)
    sport_type = data.get('sport_type', '').strip()
    technique_action = data.get('technique_action', '').strip()
    title = data.get('title', '').strip()
    content = data.get('content') or {}
    if not title or not content:
        return jsonify({'error': '缺少方案内容'})
    tp = TrainingPlan(class_id=cid, sport_type=sport_type, technique_action=technique_action,
                      title=title, content=json.dumps(content, ensure_ascii=False))
    db.session.add(tp)
    db.session.commit()
    return jsonify({'ok': True, 'id': tp.id})


@modules_bp.route('/action/training/pdf', methods=['POST'])
@login_required
def action_training_pdf():
    """专项训练方案 PDF（打印友好 HTML）。"""
    raw = request.form.get('data', '')
    try:
        data = json.loads(raw) if raw else {}
    except Exception:
        data = {}
    plan = data.get('plan') or {}
    return render_template('action/training_pdf.html', plan=plan,
                           title=str(data.get('title') or '专项训练方案'),
                           subtitle=str(data.get('subtitle') or ''))


# ============ 6. 学生动作成长档案 ============

@modules_bp.route('/action/profile')
@login_required
def action_profile():
    classes = _all_classes()
    students = Student.query.order_by(Student.class_id, Student.student_no).all()
    class_map = {c.id: c.name for c in classes}
    return render_template('action/profile.html', active='action_profile',
                           classes=classes, students=students, class_map=class_map)


@modules_bp.route('/action/profile/data')
@login_required
def action_profile_data():
    """学生成长档案数据：时间线 + 关节角度趋势 + 错误演变。"""
    student_id = request.args.get('student_id', type=int)
    if not student_id:
        return jsonify({'error': '请选择学生'})
    s = db.session.get(Student, student_id)
    if not s:
        return jsonify({'error': '学生不存在'})
    records = ActionAnalysis.query.filter_by(student_id=student_id).order_by(ActionAnalysis.id.asc()).all()
    if not records:
        return jsonify({'error': '该学生暂无动作诊断记录'})
    timeline = []
    metrics_series = {}
    for i, a in enumerate(records, 1):
        metrics = json.loads(a.metrics_data) if a.metrics_data else None
        detail = _action_detail(a)
        timeline.append({
            'id': a.id, 'index': i, 'sport_type': a.sport_type or '',
            'technique_action': a.technique_action or '未识别技术动作',
            'created_at': detail['created_at'], 'status': a.status or 'success',
            'result': a.result or '',
        })
        if metrics:
            for joint, m in metrics.items():
                metrics_series.setdefault(joint, []).append({
                    'index': i, 'avg_angle': m.get('avg_angle'), 'dev': m.get('dev'),
                    'std_angle': m.get('std_angle'),
                })
    return jsonify({'student': s.name, 'total': len(records),
                    'timeline': timeline, 'metrics_series': metrics_series})


# ============ 公共接口 ============

@modules_bp.route('/action/<int:aid>/detail')
@login_required
def action_detail(aid):
    """单条诊断详情（供弹窗渲染）。"""
    a = db.get_or_404(ActionAnalysis, aid)
    return jsonify(_action_detail(a))


@modules_bp.route('/action/<int:aid>/delete', methods=['POST'])
@login_required
def action_delete(aid):
    a = db.get_or_404(ActionAnalysis, aid)
    cid, sport_type = a.class_id, a.sport_type
    # 删除存储的视频文件
    if a.video_path:
        vp = os.path.join(UPLOAD_DIR, os.path.basename(a.video_path))
        if os.path.exists(vp):
            try:
                os.remove(vp)
            except Exception:
                pass
    # 删除抽取的帧文件
    if a.frames:
        try:
            for fp in json.loads(a.frames):
                if os.path.exists(fp):
                    os.remove(fp)
        except Exception:
            pass
    db.session.delete(a)
    db.session.commit()
    flash('已删除动作诊断记录', 'success')
    return redirect(request.referrer or url_for('modules.action_history', class_id=cid, sport_type=sport_type))


@modules_bp.route('/action/mediapipe-status')
@login_required
def action_mediapipe_status():
    """环境检测接口：返回 MediaPipe 是否已安装。"""
    return jsonify({'mediapipe_installed': MP_AVAILABLE})
